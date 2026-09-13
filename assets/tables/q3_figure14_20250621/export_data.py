"""Export Figure 14 data from the official workbook; replay, do not re-optimize."""
from __future__ import annotations

import argparse
import hashlib
import json
import subprocess
import sys
from datetime import date, timedelta
from pathlib import Path

sys.dont_write_bytecode = True
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--model-repo", type=Path, required=True)
    parser.add_argument("--date", default="2025-06-21")
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    args = parser.parse_args()
    repo = args.model_repo.resolve()
    target = date.fromisoformat(args.date)
    sys.path.insert(0, str(repo / "code/Q3"))
    import config
    from export_results import template_header_slot
    from load_data import load_prices, load_year_actuals
    from simulate import simulate_day, validate_actual

    manifest_path = repo / "results/Q3/run_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    if manifest["metrics"]["official_winner"] != "LA" or manifest["reserve_gamma"] != 0:
        raise ValueError("This export requires the official LA policy with reserve_gamma=0.")
    for key, path in (("attachment1", config.ATTACHMENT1_XLSX),
                      ("attachment2", config.ATTACHMENT2_XLSX)):
        if sha256(path) != manifest["inputs"][key]:
            raise ValueError(f"Frozen input hash mismatch: {key}")

    workbook_path = repo / "results/Q3/result3.xlsx"
    wb = load_workbook(workbook_path, read_only=True, data_only=True)
    dt = config.DT_HOURS
    checks: list[dict] = []

    def compare(name: str, actual, expected, atol: float = 1e-5) -> None:
        a, b = np.asarray(actual, dtype=float), np.asarray(expected, dtype=float)
        error = float(np.max(np.abs(a - b)))
        ok = bool(np.allclose(a, b, rtol=0, atol=atol))
        checks.append({"check": name, "passed": ok, "max_abs_error": error, "tolerance": atol})
        if not ok:
            raise ValueError(f"{name}: max error {error} exceeds {atol}")

    def extract_purchase(sheet: str):
        rows = list(wb[sheet].values)
        headers = [str(x) for x in rows[0][1:145]]
        slot_map = [template_header_slot(h) for h in headers]
        lookup = {str(pd.Timestamp(row[0]).date()): (i + 1, row)
                  for i, row in enumerate(rows[1:], start=1) if row[0] is not None}
        current_excel_row, current = lookup[target.isoformat()]
        previous_excel_row, previous = lookup[(target - timedelta(days=1)).isoformat()]
        values = np.full(144, np.nan)
        cells = [""] * 144
        for c, (day_offset, slot) in enumerate(slot_map, start=2):
            if day_offset == 0:
                values[slot] = float(current[c - 1])
                cells[slot] = f"{sheet}!{get_column_letter(c)}{current_excel_row}"
            elif (day_offset, slot) == (1, 0):
                values[0] = float(previous[c - 1])
                cells[0] = f"{sheet}!{get_column_letter(c)}{previous_excel_row}"
            else:
                raise ValueError(f"Unexpected header: {headers[c-2]}")
        if not np.isfinite(values).all():
            raise ValueError("Incomplete calendar-day purchase series")
        return values, cells, float(current[146])

    plan, plan_cells, official_plan_cost = extract_purchase("计划购电量")
    adjusted, adjusted_cells, official_total_cost = extract_purchase("调整购电量")
    storage_rows = []
    current_date = None
    for row in list(wb["充放电量"].values)[1:]:
        if row[0] is not None:
            current_date = str(pd.Timestamp(row[0]).date())
        if current_date == target.isoformat():
            storage_rows.append(row)
    if len(storage_rows) != 6:
        raise ValueError("Expected six four-hour storage records")
    soc0, soc24 = float(storage_rows[0][5]), float(storage_rows[1][5])
    official_emergency = 0.0
    current_date = None
    for row in list(wb["紧急购电量"].values)[1:]:
        if row[0] is not None:
            current_date = str(pd.Timestamp(row[0]).date())
        if current_date == target.isoformat():
            official_emergency += float(row[2] or 0.0)
    wb.close()

    prices = load_prices()
    year = load_year_actuals(jan1_load=float(prices.iloc[0].typical_load_kw),
                            jan1_pv=float(prices.iloc[0].typical_pv_kw))
    dates = pd.to_datetime(year["dates"]).dt.strftime("%Y-%m-%d").tolist()
    day_index = dates.index(target.isoformat())
    load_kw, pv_kw = year["load_kw"][day_index], year["pv_kw"][day_index]
    price = prices["price"].to_numpy(dtype=float)
    actual = simulate_day(price, load_kw * dt, pv_kw * dt, adjusted, soc0)
    errors = validate_actual(actual, adjusted, load_kw * dt, pv_kw * dt)
    if errors:
        raise ValueError(errors)
    plus, minus = np.maximum(adjusted-plan, 0), np.maximum(plan-adjusted, 0)
    plan_cost = price * plan
    deviation_cost = price * (1.5 * plus - 0.5 * minus)
    emergency_cost = 5 * price * actual["emergency_kwh"]
    total_cost = plan_cost + deviation_cost + emergency_cost
    compare("calendar_day_plan_cost_yuan", plan_cost.sum(), official_plan_cost)
    compare("calendar_day_total_cost_yuan", total_cost.sum(), official_total_cost)
    compare("day_end_soc_kwh", actual["soc24_kwh"], soc24)
    compare("emergency_purchase_kwh", actual["emergency_kwh"].sum(), official_emergency)
    compare("locked_purchase_before_06h_kwh", adjusted[:36], plan[:36])
    block_records = []
    for i, row in enumerate(storage_rows):
        charge = float(actual["charge_kwh"][24*i:24*(i+1)].sum())
        discharge = float(actual["discharge_kwh"][24*i:24*(i+1)].sum())
        compare(f"block_{i}_charge_kwh", charge, row[2])
        compare(f"block_{i}_discharge_kwh", discharge, row[3])
        block_records.append({"period": row[1], "official_charge_kwh": float(row[2]),
                              "replayed_charge_kwh": charge,
                              "official_discharge_kwh": float(row[3]),
                              "replayed_discharge_kwh": discharge})
    soc_edges = np.r_[soc0, actual["soc_end_kwh"]]
    compare("energy_balance_kwh", adjusted + pv_kw*dt + actual["discharge_kwh"]
            + actual["emergency_kwh"] - load_kw*dt - actual["charge_kwh"]
            - actual["curtail_kwh"], np.zeros(144))
    compare("soc_dynamics_kwh", np.diff(soc_edges), config.ETA_CHARGE*actual["charge_kwh"]
            - actual["discharge_kwh"]/config.ETA_DISCHARGE)

    def clock(minute: int) -> str:
        return f"{minute//60:02d}:{minute%60:02d}"

    source_direct = "附件/官方结果直接提取或单位换算"
    source_replay = "按官方锁定购电量及原执行规则复算；已与4小时汇总及日末储电量核对"
    columns = [
        ("slot", np.arange(144), "0起始的日历时段编号", "—", source_direct),
        ("time_start", [clock(i*10) for i in range(144)], "时段起点", "HH:MM", source_direct),
        ("time_end", [clock((i+1)*10) for i in range(144)], "时段终点", "HH:MM", source_direct),
        ("hour_start", np.arange(144)*dt, "绘图横坐标：时段起点", "h", source_direct),
        ("hour_end", np.arange(1,145)*dt, "时段终点横坐标", "h", source_direct),
        ("load_kw", load_kw, "实测小区负载功率", "kW", source_direct),
        ("pv_kw", pv_kw, "实测光伏功率", "kW", source_direct),
        ("net_load_kw", load_kw-pv_kw, "净负载=负载-光伏，可为负", "kW", source_direct),
        ("price_yuan_per_kwh", price, "本时段交易电价", "元/kWh", source_direct),
        ("planned_purchase_kwh", plan, "0:00原始计划购电量", "kWh/时段", source_direct),
        ("adjusted_purchase_kwh", adjusted, "最终执行的调整购电量", "kWh/时段", source_direct),
        ("planned_purchase_kw", plan/dt, "计划购电平均功率", "kW", source_direct),
        ("adjusted_purchase_kw", adjusted/dt, "调整购电平均功率", "kW", source_direct),
        ("charge_kwh", actual["charge_kwh"], "微网侧充电量，流入电池前", "kWh/时段", source_replay),
        ("discharge_kwh", actual["discharge_kwh"], "微网侧放电量，电池输出后", "kWh/时段", source_replay),
        ("charge_kw", actual["charge_kwh"]/dt, "微网侧充电平均功率", "kW", source_replay),
        ("discharge_kw", actual["discharge_kwh"]/dt, "微网侧放电平均功率", "kW", source_replay),
        ("storage_net_output_kw", (actual["discharge_kwh"]-actual["charge_kwh"])/dt,
         "储能对微网净输出：正为放电，负为充电", "kW", source_replay),
        ("emergency_purchase_kwh", actual["emergency_kwh"], "紧急购电量", "kWh/时段", source_replay),
        ("emergency_purchase_kw", actual["emergency_kwh"]/dt, "紧急购电平均功率", "kW", source_replay),
        ("curtailed_pv_kwh", actual["curtail_kwh"], "弃光电量", "kWh/时段", source_replay),
        ("curtailed_pv_kw", actual["curtail_kwh"]/dt, "弃光平均功率", "kW", source_replay),
        ("soc_start_kwh", soc_edges[:-1], "时段起点储电量", "kWh", source_replay),
        ("soc_end_kwh", soc_edges[1:], "时段终点储电量", "kWh", source_replay),
        ("soc_min_kwh", np.full(144, config.E_MIN_KWH), "储电量下限", "kWh", source_direct),
        ("soc_max_kwh", np.full(144, config.E_MAX_KWH), "储电量上限", "kWh", source_direct),
        ("plan_cost_yuan", plan_cost, "原始计划购电费", "元/时段", source_direct),
        ("adjustment_cost_yuan", deviation_cost, "调整增量费用，退费可为负", "元/时段", source_direct),
        ("emergency_cost_yuan", emergency_cost, "紧急购电费", "元/时段", source_replay),
        ("total_cost_yuan", total_cost, "计划+调整增量+紧急购电费", "元/时段", source_replay),
        ("plan_source_cell", plan_cells, "计划购电在官方工作簿中的来源单元格", "—", source_direct),
        ("adjusted_source_cell", adjusted_cells, "调整购电在官方工作簿中的来源单元格", "—", source_direct),
    ]
    slots = pd.DataFrame({c[0]: c[1] for c in columns})
    dictionary = pd.DataFrame([{"field": c[0], "meaning": c[2], "unit": c[3], "provenance": c[4]}
                               for c in columns])
    edges = pd.DataFrame({"time": [clock(i*10) for i in range(145)], "hour": np.arange(145)*dt,
                          "soc_kwh": soc_edges, "soc_min_kwh": config.E_MIN_KWH,
                          "soc_max_kwh": config.E_MAX_KWH})
    summary = {
        "date": target.isoformat(), "policy": "LA", "q_lock": manifest["q_lock"],
        "duration_hours": 24, "slot_hours": dt, "slot_count": 144, "soc_point_count": 145,
        "planned_purchase_kwh": float(plan.sum()), "adjusted_purchase_kwh": float(adjusted.sum()),
        "adjusted_minus_plan_kwh": float(adjusted.sum()-plan.sum()),
        "plan_cost_yuan": float(plan_cost.sum()), "adjustment_cost_yuan": float(deviation_cost.sum()),
        "emergency_purchase_kwh": float(actual["emergency_kwh"].sum()),
        "emergency_cost_yuan": float(emergency_cost.sum()), "total_cost_yuan": float(total_cost.sum()),
        "soc_start_kwh": soc0, "soc_end_kwh": float(soc_edges[-1]),
        "soc_min_observed_kwh": float(soc_edges.min()), "soc_max_observed_kwh": float(soc_edges.max()),
        "charge_total_kwh": float(actual["charge_kwh"].sum()),
        "discharge_total_kwh": float(actual["discharge_kwh"].sum()),
        "max_charge_kw": float(actual["charge_kwh"].max()/dt),
        "max_discharge_kw": float(actual["discharge_kwh"].max()/dt),
        "max_adjacent_soc_change_kwh": float(np.abs(np.diff(soc_edges)).max()),
        "validation_passed": True,
    }
    source_paths = [workbook_path, manifest_path, config.ATTACHMENT1_XLSX, config.ATTACHMENT2_XLSX,
                    repo/"code/Q3/config.py", repo/"code/Q3/load_data.py", repo/"code/Q3/simulate.py",
                    repo/"code/Q3/export_results.py", repo/"code/common/time_slots.py"]
    sources = [{"path": p.relative_to(repo).as_posix(), "sha256": sha256(p)} for p in source_paths]
    provenance = {
        "model_commit": subprocess.check_output(["git", "-C", str(repo), "rev-parse", "HEAD"], text=True).strip(),
        "date": target.isoformat(), "sources": sources,
        "time_alignment": "calendar [00:00,24:00); slot 0 from previous template row's final column",
        "reconstructed_fields": [c[0] for c in columns if c[4] == source_replay],
        "method": "Deterministic execution replay from official final purchases and actual load/PV; no optimization rerun.",
        "ai_assistance": "Codex prepared this export and checked source alignment and numerical consistency.",
        "checks": checks,
    }
    out = args.output_dir.resolve()
    out.mkdir(parents=True, exist_ok=True)
    for name, frame in (("plot_data_10min.csv", slots), ("soc_145_points.csv", edges),
                        ("data_dictionary.csv", dictionary), ("four_hour_validation.csv", pd.DataFrame(block_records))):
        frame.to_csv(out/name, index=False, encoding="utf-8-sig")
    (out/"summary.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    (out/"provenance.json").write_text(json.dumps(provenance, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    book_path = out / f"figure14_data_{target.strftime('%Y%m%d')}.xlsx"
    with pd.ExcelWriter(book_path, engine="openpyxl") as writer:
        for name, frame in (("plot_10min", slots), ("soc_145_points", edges),
                            ("summary", pd.DataFrame(summary.items(), columns=["metric", "value"])),
                            ("data_dictionary", dictionary), ("four_hour_check", pd.DataFrame(block_records)),
                            ("validation", pd.DataFrame(checks)), ("sources", pd.DataFrame(sources))):
            frame.to_excel(writer, sheet_name=name, index=False)
            ws = writer.sheets[name]
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            ws.row_dimensions[1].height = 30
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="24476B")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            for col in ws.columns:
                letter = col[0].column_letter
                width = max(len(str(c.value or "")) for c in list(col)[:20]) + 2
                ws.column_dimensions[letter].width = min(65, max(15, width))
                for cell in col[1:]:
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        cell.number_format = "0.000000"
    check_wb = load_workbook(book_path, read_only=True, data_only=True)
    assert check_wb["plot_10min"].max_row == 145
    assert check_wb["soc_145_points"].max_row == 146
    check_wb.close()
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(f"Validated {len(checks)} comparisons. Export: {book_path}")


if __name__ == "__main__":
    main()
